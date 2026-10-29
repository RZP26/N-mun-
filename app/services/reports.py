from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

def make_excel(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Nəticələr'
    headers = ['Tələbə kodu', 'Tələbə', 'Qiymətləndirmə', 'Toplanan bal', 'Maksimum bal', 'Faiz', 'Düzgün', 'Səhv', 'Boş', 'Yoxlama tələb olunur']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.alignment = Alignment(wrap_text=True)
    for row in rows:
        ws.append([row.get('student_no',''), row.get('student_name','Naməlum'), row.get('assessment_title',''), row.get('score',0), row.get('max_score',0), row.get('percent',0), row.get('correct_count',0), row.get('wrong_count',0), row.get('blank_count',0), row.get('needs_review','')])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col, width in {'A':16,'B':28,'C':28,'D':14,'E':14,'F':12,'G':10,'H':10,'I':10,'J':28}.items(): ws.column_dimensions[col].width = width
    out = BytesIO(); wb.save(out); return out.getvalue()

def make_pdf(rows: list[dict]) -> bytes:
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [Paragraph('Peşə təhsili: qiymətləndirmə hesabatı', styles['Title']), Spacer(1, 12)]
    data = [['Kod','Tələbə','Qiymətləndirmə','Bal','Maks.','Faiz','Düzgün','Səhv','Boş','Yoxlama']]
    for r in rows:
        data.append([str(r.get('student_no','')), str(r.get('student_name','Naməlum'))[:35], str(r.get('assessment_title',''))[:28], str(r.get('score',0)), str(r.get('max_score',0)), f"{r.get('percent',0)}%", str(r.get('correct_count',0)), str(r.get('wrong_count',0)), str(r.get('blank_count',0)), str(r.get('needs_review',''))[:18]])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#17365D')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),0.4,colors.grey),('FONTSIZE',(0,0),(-1,-1),8),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white, colors.HexColor('#F2F6FA')])]))
    story.append(table); doc.build(story); return out.getvalue()
